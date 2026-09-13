"""Shared test fixtures.

Resume fixtures are generated at runtime rather than committed as binaries, so
the repo carries no sample PDFs and the fixtures never drift from the text the
assertions expect.
"""

import os

# Must run before app.core.config is imported anywhere: the settings object is
# cached, so a later change would have no effect. Disabling the semantic tier
# keeps the suite fast and stops CI downloading a 90MB model. The tier itself is
# covered by tests/test_semantic.py, which is opt-in.
os.environ.setdefault("ENABLE_SEMANTIC", "false")

from pathlib import Path

import pytest
from docx import Document
from reportlab.lib.pagesizes import A4
from reportlab.pdfgen import canvas

SAMPLE_RESUME_LINES: list[str] = [
    "Priya Sharma",
    "priya.sharma@example.com | +91 98765 43210 | linkedin.com/in/priyasharma",
    "",
    "PROFESSIONAL SUMMARY",
    "Final-year computer science student with internship experience in backend",
    "development and applied machine learning.",
    "",
    "TECHNICAL SKILLS",
    "Languages: Python, Java, C++, C#, JavaScript, SQL",
    "Frameworks: React, Node.js, ASP.NET, FastAPI",
    "Practices: CI/CD, Agile, Unit Testing",
    "",
    "WORK EXPERIENCE",
    "Backend Intern, Nimbus Labs (2024)",
    "Built REST APIs serving 20k requests per day and cut p95 latency by 40%.",
    "",
    "PROJECTS",
    "Resume Analyzer - NLP pipeline using scikit-learn and TF-IDF.",
    "",
    "EDUCATION",
    "B.Tech Computer Science, 2022 - 2026",
    "",
    "CERTIFICATIONS",
    "AWS Certified Cloud Practitioner, 2025",
]


@pytest.fixture(scope="session")
def sample_lines() -> list[str]:
    """Return the text lines every generated fixture resume contains."""
    return SAMPLE_RESUME_LINES


@pytest.fixture(scope="session")
def sample_docx(tmp_path_factory: pytest.TempPathFactory) -> Path:
    """Build a .docx resume on disk and return its path."""
    path = tmp_path_factory.mktemp("fixtures") / "resume.docx"
    document = Document()
    for line in SAMPLE_RESUME_LINES:
        document.add_paragraph(line)
    document.save(str(path))
    return path


@pytest.fixture(scope="session")
def sample_docx_with_table(tmp_path_factory: pytest.TempPathFactory) -> Path:
    """Build a .docx whose skills live in a table, and return its path.

    Many resume templates do this; paragraph-only parsing would miss every skill.
    """
    path = tmp_path_factory.mktemp("fixtures") / "resume_table.docx"
    document = Document()
    document.add_paragraph("TECHNICAL SKILLS")
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "Languages"
    table.cell(0, 1).text = "Python, C++, JavaScript"
    table.cell(1, 0).text = "Practices"
    table.cell(1, 1).text = "CI/CD, Docker, Kubernetes"
    # Padding so the result clears MIN_MEANINGFUL_CHARS the way a real resume does.
    document.add_paragraph("WORK EXPERIENCE")
    document.add_paragraph("Backend Intern at Nimbus Labs, built REST APIs and "
                           "automated the deployment pipeline end to end.")
    document.save(str(path))
    return path


@pytest.fixture(scope="session")
def sample_pdf(tmp_path_factory: pytest.TempPathFactory) -> Path:
    """Build a text-based .pdf resume on disk and return its path."""
    path = tmp_path_factory.mktemp("fixtures") / "resume.pdf"
    page = canvas.Canvas(str(path), pagesize=A4)
    page.setFont("Helvetica", 11)
    y = 800
    for line in SAMPLE_RESUME_LINES:
        page.drawString(50, y, line)
        y -= 16
    page.save()
    return path


@pytest.fixture(scope="session")
def blank_pdf(tmp_path_factory: pytest.TempPathFactory) -> Path:
    """Build a PDF with no text layer, standing in for a scanned resume."""
    path = tmp_path_factory.mktemp("fixtures") / "blank.pdf"
    page = canvas.Canvas(str(path), pagesize=A4)
    page.showPage()
    page.save()
    return path
